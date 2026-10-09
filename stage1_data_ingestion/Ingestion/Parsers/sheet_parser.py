from typing import Any

import pandas as pd
from stage1_data_ingestion.Ingestion.Parsers.document_parser import DocumentParser
from stage1_data_ingestion.Ingestion.models import Document
import numpy as np
import openpyxl
from openpyxl.utils.cell import range_boundaries
import pandas as pd
from scipy.ndimage import label

class SheetParser (DocumentParser):
    def __init__(self):
        pass

    def parse_cell_range(self, cell_range):
            '''Converts a range string like 'A1:C10' to 1-indexed boundaries (min_col, min_row, max_col, max_row).'''
            return range_boundaries(cell_range)

    def scan_workbook(self, file_path):
        wb = openpyxl.load_workbook(file_path, data_only=True)
        all_tables = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            max_row, max_col = ws.max_row, ws.max_column

            if max_row == 0 or max_col == 0:
                continue

            # Matrix to track non-empty cells (1 = occupied, 0 = empty)
            grid = np.zeros((max_row, max_col), dtype=int)
            for r_idx, row in enumerate(
                ws.iter_rows(values_only=True), start=1
            ):
                for c_idx, val in enumerate(row, start=1):
                    if val is not None and str(val).strip() != "":
                        grid[r_idx - 1, c_idx - 1] = 1

            # Matrix to mark cells already claimed by formal tables
            claimed_mask = np.zeros((max_row, max_col), dtype=bool)

            # -------------------------------------------------------------
            # STEP 1: Process Formal Excel Tables (ListObjects)
            # -------------------------------------------------------------
            for table_name, cell_range in ws.tables.items():
                min_col, min_row, max_col, max_row = self.parse_cell_range(
                    cell_range
                )

                # Mark this region as claimed (0-indexed array slices)
                claimed_mask[
                    min_row - 1 : max_row, min_col - 1 : max_col
                ] = True

                # Extract cell values
                data = [
                    [ws.cell(row=r, column=c).value for c in range(min_col, max_col + 1)]
                    for r in range(min_row, max_row + 1)
                ]

                # Build DataFrame
                headers = data[0]
                rows = data[1:]
                df = pd.DataFrame(rows, columns=headers)

                all_tables.append(
                    {
                        "sheet": sheet_name,
                        "type": "Formal Table",
                        "name": table_name,
                        "range": cell_range,
                        "dataframe": df,
                    }
                )

            # -------------------------------------------------------------
            # STEP 2: Process Informal Data Blocks (Unregistered Regions)
            # -------------------------------------------------------------
            # Zero out grid cells that belong to formal tables
            grid[claimed_mask] = 0

            # Find remaining contiguous data regions
            labeled_array, num_features = label(grid)

            for i in range(1, num_features + 1):
                rows, cols = np.where(labeled_array == i)

                # Bounds of the implicit table
                min_r, max_r = rows.min() + 1, rows.max() + 1
                min_c, max_c = cols.min() + 1, cols.max() + 1

                # Extract block content
                block_data = [
                    [ws.cell(row=r, column=c).value for c in range(min_c, max_c + 1)]
                    for r in range(min_r, max_r + 1)
                ]

                # Assume first row of the implicit block is the header
                headers = block_data[0]
                rows = block_data[1:] if len(block_data) > 1 else []
                df = pd.DataFrame(rows, columns=headers)
                #print(f'Row: {rows}, df:{df}')
                range_str = f"R{min_r}C{min_c}:R{max_r}C{max_c}"

                all_tables.append(
                    {
                        "sheet": sheet_name,
                        "type": "Implicit Region",
                        "name": f"Implicit_{range_str}",
                        "range": range_str,
                        "dataframe": df,
                    }
                )

        return all_tables
    
        
       
    def parse(self,file_path) -> Document:
        '''df = pd.read_excel(file_path)

        # 2. Inspect the data structure
        print("--- First 5 Rows ---")
        print(df.head())

        print("\n--- Column Names & Data Types ---")
        print(df.info())'''
        # -------------------------------------------------------------
        # Usage Example
        # -------------------------------------------------------------
        results = self.scan_workbook(file_path)
        document:Document = Document()
        document.tables_list = results

        for item in results:
            #print(f"\nSheet: {item['sheet']} | Type: {item['type']} | Range: {item['range']}")
            print('item contents: ',item["dataframe"], ' item data type', type(item))

        return document
    

        